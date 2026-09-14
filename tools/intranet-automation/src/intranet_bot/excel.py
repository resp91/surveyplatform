"""지정된 엑셀 파일에서 입력할 내용을 읽어온다."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterator

from .config import ExcelConfig
from .errors import ExcelError


@dataclass
class ExcelRow:
    """엑셀 한 행. row_number 는 엑셀에서 보이는 실제 행 번호(1-based)."""

    row_number: int
    values: dict[str, Any]

    def is_empty(self) -> bool:
        return all(v in (None, "") for v in self.values.values())


class ExcelSource:
    """시나리오의 excel 설정대로 시트를 열고 행을 돌려준다."""

    def __init__(self, config: ExcelConfig, path: Path):
        self.config = config
        self.path = path
        self._rows: list[ExcelRow] | None = None

    def load(self) -> list[ExcelRow]:
        if self._rows is not None:
            return self._rows

        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - 설치 안내용
            raise ExcelError("openpyxl 이 필요합니다. 'pip install -r requirements.txt' 를 실행하세요.") from exc

        if not self.path.is_file():
            raise ExcelError(f"엑셀 파일을 찾을 수 없습니다: {self.path}")

        # data_only=True: 수식 대신 엑셀이 마지막으로 계산해 저장한 값을 읽는다.
        workbook = load_workbook(self.path, data_only=True, read_only=True)
        try:
            if self.config.sheet:
                if self.config.sheet not in workbook.sheetnames:
                    raise ExcelError(
                        f"'{self.config.sheet}' 시트가 없습니다. 있는 시트: {', '.join(workbook.sheetnames)}"
                    )
                sheet = workbook[self.config.sheet]
            else:
                sheet = workbook[workbook.sheetnames[0]]

            grid = [list(r) for r in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()

        header_index = self.config.header_row - 1
        if header_index < 0 or header_index >= len(grid):
            raise ExcelError(
                f"header_row({self.config.header_row})가 시트 범위를 벗어났습니다. "
                f"시트 행 수: {len(grid)}"
            )

        headers = [_clean_header(h) for h in grid[header_index]]
        column_index = _map_columns(self.config.columns, headers)

        last_row = self.config.end_row or len(grid)
        rows: list[ExcelRow] = []
        for row_number in range(self.config.start_row, min(last_row, len(grid)) + 1):
            raw = grid[row_number - 1]
            values = {
                name: (raw[idx] if idx < len(raw) else None)
                for name, idx in column_index.items()
            }
            row = ExcelRow(row_number=row_number, values=values)
            if self.config.skip_empty_rows and row.is_empty():
                continue
            rows.append(row)

        self._rows = rows
        return rows

    def __iter__(self) -> Iterator[ExcelRow]:
        return iter(self.load())


def _clean_header(value: Any) -> str:
    """헤더 셀의 공백/줄바꿈을 정리한다. 사내 양식은 헤더에 줄바꿈이 흔하다."""
    if value is None:
        return ""
    return " ".join(str(value).split())


def _map_columns(columns: dict[str, str], headers: list[str]) -> dict[str, int]:
    """'엑셀 헤더명 -> 변수명' 설정을 '변수명 -> 열 인덱스'로 바꾼다.

    헤더명 대신 'A', 'B' 같은 열 문자로도 지정할 수 있다.
    """
    if not columns:
        # 설정이 없으면 헤더 이름을 그대로 변수명으로 쓴다.
        return {h: i for i, h in enumerate(headers) if h}

    mapping: dict[str, int] = {}
    for source, var_name in columns.items():
        key = _clean_header(source)
        if key in headers:
            mapping[var_name] = headers.index(key)
            continue
        letter_index = _column_letter_to_index(key)
        if letter_index is not None:
            mapping[var_name] = letter_index
            continue
        available = ", ".join(h for h in headers if h) or "(헤더 없음)"
        raise ExcelError(f"엑셀에 '{source}' 열이 없습니다. 있는 열: {available}")
    return mapping


def _column_letter_to_index(text: str) -> int | None:
    """'A' -> 0, 'B' -> 1, 'AA' -> 26. 열 문자가 아니면 None."""
    if not text or not text.isalpha() or not text.isascii():
        return None
    index = 0
    for char in text.upper():
        index = index * 26 + (ord(char) - ord("A") + 1)
    return index - 1
